import socket
import sys
from openpyxl import load_workbook
from openpyxl import Workbook
from datetime import datetime
import time
import datetime

# measure ping time over tcp
def tcp_ping(ip, port):
    try:
        start_time = time.perf_counter()
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.settimeout(5)
        s.connect((ip, port))
        s.shutdown(socket.SHUT_RD)
        
        if s.getsockopt(socket.SOL_SOCKET, socket.SO_ERROR) == 0:
            end_time = time.perf_counter()
            ping_time = end_time - start_time
            service_name = socket.getservbyport(port)
            s.close()
            time.sleep(1) 
            print(f"PORT {port}/tcp - STATE: open - SERVICE: {service_name}")
            print(f"Host is up - {ping_time*1000:.3f}ms latency")
            return f"{ping_time*1000:.3f}"
        else:
            print("TCP handshake failed")
            s.close()
            return 'N/A'
    except (socket.timeout, socket.error) as e:
        print("An error occurred while connecting to the server:", e)
        sys.exit(1)
        return 'N/A'

def log_to_excel(ip, port, protocol, ping_time_log, repeats):
    try:
        wb = load_workbook('network_ping_log.xlsx')
    except FileNotFoundError:
        wb = Workbook()
    
    timestamp = datetime.datetime.now().strftime('%d-%m-%Y-%H%M%S%f')[:-3]
    
    sheet_name = f'{timestamp}'
    if sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
    else:
        sheet = wb.create_sheet(sheet_name)
    
    sheet.append(['Attempts','IP Address', 'Port', 'Protocol', 'Ping Time(ms)'])
    for i in range(repeats):
        sheet.append([i+1, ip, port, protocol, ping_time_log[i]])

    # Calculate the average time
    valid_ping_times = [float(time) for time in ping_time_log if time != 'N/A']
    if valid_ping_times:
        average_ping_time = sum(valid_ping_times) / len(valid_ping_times)
        print(f"Average ping time: {average_ping_time:.3f} ms")
    else:
        print("No successful pings.")
    sheet.append(['','','','Average ping time: ', f"{average_ping_time:.3f} ms"])

    # Save the workbook
    wb.save('network_ping_log.xlsx')
    time.sleep(2)
    print("network_ping_log.xlsx successfully created.")

if __name__ == "__main__":
    ip_address = input("Enter the IP address to ping: ")
    port = int(input("Enter the port number: "))
    repeat = int(input("Enter the number of times to ping: "))

    ping_time_log = []
    for i in range(repeat):
        print(f"Attempt {i + 1}: Pinging {ip_address}:{port} using TCP... ")
        ping_time = tcp_ping(ip_address, port)
        ping_time_log.append(ping_time)
    log_to_excel(ip_address, port, 'TCP', ping_time_log, repeat)
